"""Extrae texto de archivos subidos: PDF, Word (.docx), OpenDocument (.odt), RTF, TXT y similares."""
import io
import re
import zipfile
from pathlib import PurePath

MAX_CHARS = 30000
SOPORTADOS = {".pdf", ".docx", ".odt", ".rtf", ".txt", ".md", ".csv"}


class ErrorArchivo(Exception):
    """Mensaje pensado para mostrarse tal cual al usuario."""


def _pdf(datos):
    from pypdf import PdfReader
    try:
        lector = PdfReader(io.BytesIO(datos))
        if lector.is_encrypted:
            raise ErrorArchivo("El PDF está protegido con contraseña. Quita la protección y vuelve a subirlo.")
        texto = "\n".join((p.extract_text() or "") for p in lector.pages)
    except ErrorArchivo:
        raise
    except Exception:
        raise ErrorArchivo("No se pudo leer el PDF: parece dañado o no es un PDF válido.")
    if len(texto.strip()) < 30:
        raise ErrorArchivo("El PDF no tiene texto seleccionable (probablemente es un escaneo). "
                           "Esta versión no aplica OCR: sube una versión con texto o pega el contenido.")
    return texto


def _docx(datos):
    import docx
    try:
        d = docx.Document(io.BytesIO(datos))
    except Exception:
        raise ErrorArchivo("No se pudo leer el archivo Word: parece dañado o no es un .docx válido.")
    partes = [p.text for p in d.paragraphs]
    for tabla in d.tables:
        for fila in tabla.rows:
            partes.append(" ".join(c.text for c in fila.cells))
    return "\n".join(partes)


def _odt(datos):
    try:
        with zipfile.ZipFile(io.BytesIO(datos)) as z:
            xml = z.read("content.xml").decode("utf-8")
    except Exception:
        raise ErrorArchivo("No se pudo leer el archivo .odt: parece dañado.")
    xml = re.sub(r"</text:(p|h)>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def _rtf(datos):
    t = datos.decode("latin-1")
    t = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "replace"), t)
    t = re.sub(r"\\pard?", "\n", t)
    t = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", t)
    return re.sub(r"[{}]", "", t)


def _plano(datos):
    for cod in ("utf-8-sig", "cp1252"):
        try:
            return datos.decode(cod)
        except UnicodeDecodeError:
            pass
    raise ErrorArchivo("No se pudo decodificar el archivo de texto.")


def extraer_texto(nombre, datos):
    """Devuelve el texto limpio del archivo o lanza ErrorArchivo con un mensaje claro."""
    ext = PurePath(nombre).suffix.lower()
    if ext == ".doc":
        raise ErrorArchivo("El formato .doc (Word antiguo) no se puede leer. Ábrelo y guárdalo como .docx o PDF.")
    if ext not in SOPORTADOS:
        raise ErrorArchivo(f"Formato '{ext or 'sin extensión'}' no soportado. Usa PDF, Word (.docx), .odt, .rtf o .txt.")
    if not datos:
        raise ErrorArchivo("El archivo está vacío.")
    lector = {".pdf": _pdf, ".docx": _docx, ".odt": _odt, ".rtf": _rtf}.get(ext, _plano)
    texto = re.sub(r"[ \t]+", " ", lector(datos))
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    if len(texto) < 40:
        raise ErrorArchivo("El archivo tiene muy poco texto (menos de 40 caracteres) para clasificarlo.")
    if len(texto) > MAX_CHARS:
        raise ErrorArchivo(f"El texto supera {MAX_CHARS:,} caracteres. Sube un fragmento más breve.".replace(",", "."))
    return texto
